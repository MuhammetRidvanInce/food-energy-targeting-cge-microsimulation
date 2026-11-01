"""Check publication assets without requesting access to confidential microdata."""
from pathlib import Path
import ast,hashlib,json,re,sys
import pandas as pd
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
def main():
    files=[p for p in ROOT.rglob('*') if p.is_file() and '.git' not in p.parts and '__pycache__' not in p.parts]
    assert not any(p.name in ['hane_cd_2024.csv','fert_cd_2024.csv','tuketim_cd_2024.csv','hbs_madde_tanim.csv'] for p in files)
    assert not any(p.suffix.lower() in ['.sav','.dta','.parquet','.pkl','.pickle','.npz','.npy','.docx'] for p in files)
    spreadsheet_count=0;csv_count=0
    for p in files:
        if p.suffix=='.py':ast.parse(p.read_text(encoding='utf-8'))
        if p.suffix=='.ipynb':
            nb=json.loads(p.read_text(encoding='utf-8'))
            for c in nb['cells']:
                if c['cell_type']=='code':
                    assert c['outputs']==[] and c['execution_count'] is None,p
                    ast.parse(''.join(c['source']))
        if p.suffix in ['.py','.json','.md','.txt','.ipynb','.csv'] and p.name!='validate_public_release.py':
            assert not re.search(r'[A-Z]:[\\/](?:Users|Python|Program Files)',p.read_text(encoding='utf-8-sig')),p
        if p.suffix=='.csv':
            df=pd.read_csv(p);csv_count+=1
            assert not set(df.columns)&{'BIRIMNO','FERTNO','YILKULGEL','EFB_OECD','HHB','FAKTOR','household_id','person_id'},p
        if p.suffix=='.xlsx':
            w=load_workbook(p,read_only=True,data_only=True);spreadsheet_count+=1
            for s in w:
                for row in s.iter_rows(values_only=True):
                    for v in row:
                        if isinstance(v,str):assert not re.search(r'[A-Z]:[\\/](?:Users|Python|Program Files)',v),(p,s.title)
            w.close()
    sys.path.insert(0,str(ROOT/'4_cge_model/codes'))
    import cge_model as cge
    bm=cge.load_benchmark();el=cge.load_elasticities()
    assert bm.sam.shape==(78,78) and el.shape[0]==64
    p=ROOT/'5_simulation/outputs/tables/main/Table1_main_policy_results.csv'
    results=pd.read_csv(p).set_index('policy')
    assert abs((results.loc['pmt','fgt1_pct']-results.loc['universal_cash','fgt1_pct'])+.387688149)<1e-7
    sup=ROOT/'publication_tables/supplementary'
    assert len(list(sup.glob('*.xlsx')))==19
    for i in range(1,8):
        assert (ROOT/f'5_simulation/outputs/figures/main/fig{i}.png').exists()
        assert (ROOT/f'5_simulation/outputs/figures/supplementary/figS{i}.png').exists()
    report={'status':'passed','csv_files':csv_count,'workbooks':spreadsheet_count,'sector_count':64,'sam_accounts':78,
            'elasticity_values':192,'supplementary_tables':19,'main_figures':7,'supplementary_figures':7,
            'raw_microdata_included':False,'notebook_outputs_included':False,'local_absolute_paths_found':False}
    print(json.dumps(report,indent=2))
    return report
if __name__=='__main__':main()
