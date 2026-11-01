"""Publish the coverage/Pareto extension using the established output folders.

Supplementary workbooks have exactly Data and Variable Descriptions worksheets,
with raw values and no formatting. Existing S1--S13 files are not rewritten.
"""
from pathlib import Path
import argparse,json,copy
import numpy as np
import pandas as pd
from openpyxl import Workbook,load_workbook
import figure_design as fd
from matplotlib.text import Text
from matplotlib.ticker import FuncFormatter
from common import OUT,PROJECT,dump,configurations

TR=PROJECT/'6_manuscript/preparation/turkish'
META={
'coverage_pct':('Persons receiving cash as a share of the represented population; weighted boundary households receive full payment in the selected state.','Percent'),
'recipients_million':('Represented persons receiving cash.','Million persons'),
'target_reached_pct':('Share of the fixed bottom-two-quintile target population receiving cash; distinct from total population coverage.','Percent'),
'initial_poor_reached_pct':('Share of persons below the initial fixed poverty line receiving cash.','Percent'),
'exclusion_pct':('Share of the fixed bottom-two-quintile target population not receiving cash.','Percent'),
'leakage_pct':('Share of recipients outside the fixed bottom-two-quintile target population; equals their cash share under uniform payments.','Percent'),
'annual_cash_per_recipient_sam_tl':('Annual equal per-person cash payment on the 2023 SAM calibration scale, not survey-year income units.','2023 calibrated TRY per person'),
'net_fiscal_cost_billion_tl':('Decline in government saving relative to the same-shock no-policy equilibrium.','Billion 2023 calibrated TRY'),
'net_cash_billion_tl':('Aggregate cash delivered to households, excluding administrative purchases.','Billion 2023 calibrated TRY'),
'fgt0_pct':('Person-weighted poverty headcount using the fixed initial real poverty line.','Percent'),
'fgt1_pct':('Person-weighted mean normalised poverty gap over the entire represented population.','Percent'),
'fgt2_pct':('Person-weighted mean squared normalised poverty gap over the entire represented population.','Percent'),
'gini':('Person-weighted Gini coefficient of real equivalent disposable income including receipt-state mixtures.','Index'),
'bottom40_ev_pct':('Change in real consumption of the initial bottom two quintiles, divided by their initial expenditure.','Percent'),
'all_ev_pct':('Aggregate change in real household consumption, divided by initial household expenditure.','Percent'),
'bottom40_loser_pct':('Share of the initial bottom-two-quintile population with real consumption below its pre-shock level.','Percent'),
'micro_macro_demand_gap_pct':('Sum of absolute differences between implied micro and CGE quintile-sector spending, divided by total CGE household spending.','Percent'),
'max_group_budget_gap_trillion_tl':('Largest absolute reconciliation error between aggregate micro budgets and a CGE household-group budget.','Trillion calibrated TRY'),
'negative_income_households':('Number of survey households with negative reconstructed income in any positive-weight receipt state.','Households'),
'fgt1_regret_pp':('FGT1 minus the numerical minimum for the main configuration and original PMT ranking.','Percentage points'),
'additional_recipients_million':('Recipients in excess of the numerically selected minimum-gap design. Negative values denote smaller coverage.','Million persons'),
'pareto_efficient_evaluated':('True when no evaluated policy has at least as much population coverage and no greater FGT1, with one strict improvement; numerical tolerance is applied.','Boolean'),
'maximum_reduction_retained_pct':('100 times (same-shock no-policy FGT1 minus policy FGT1) divided by (same-shock no-policy FGT1 minus the numerical minimum).','Percent'),
'epsilon_pp':('Allowed increase in FGT1 above the numerical minimum; delta in the manuscript epsilon-constraint formulation.','Percentage points'),
'scenario_kind':('absolute_tolerance uses an FGT1 tolerance; retention_target fixes the retained share of maximum poverty-gap reduction.','Category'),
'fgt1_upper_bound_pct':('Numerical minimum FGT1 plus the allowed tolerance.','Percent'),
'reduction_retention_requirement_pct':('Minimum share of the maximum achievable reduction relative to the no-policy shock; blank for absolute tolerances.','Percent'),
'actual_fgt1_cost_pp':('Selected policy FGT1 minus the numerical minimum.','Percentage points'),
'epsilon_slack_pp':('Permitted FGT1 upper bound minus selected policy FGT1; small negative values reflect numerical roundoff.','Percentage points'),
'configuration':('Structural scenario or PMT partition identifier. Structural settings are recorded in analysis_design.json; pmt_partition identifiers select alternative five-fold PMT random seeds.','Category'),
'seed':('Random seed defining the five shuffled household folds used for out-of-fold PMT predictions.','Integer'),
'stage':('Numerical search stage: coarse grid, local grid/refinement, dense reference, epsilon bracket/root/probe or feedback decomposition.','Category'),
'selected_choice':('minimum_gap, retain_75 or retain_50 denotes a selected design; blank identifies other evaluated candidates.','Category'),
'selected_retention_pct':('Retained-reduction requirement associated with the selected design; 100 identifies the numerical minimum.','Percent'),
'general_equilibrium_fgt1_pct':('FGT1 with solved general-equilibrium prices, regular incomes and the policy-specific cash envelope.','Percent'),
'frozen_feedback_fgt1_pct':('FGT1 holding the same cash envelope and recipients fixed while freezing prices and regular incomes at the no-policy shock.','Percent'),
'feedback_difference_pp':('General-equilibrium FGT1 minus the frozen-feedback value. This is a channel decomposition at fixed cash, not a separately balanced policy equilibrium.','Percentage points'),
'policy_a':('First fixed candidate in the paired conditional comparison. minimum_gap is the refined CES optimum; nearby_minimum is the preliminary CES candidate.','Category'),
'policy_b':('Second fixed candidate: pmt_20 or pmt_40 is exact 20% or 40% coverage; universal_cash is 100%; nearby_minimum is the preliminary CES candidate.','Category'),
'difference_pp':('Person-weighted FGT1 of policy_a minus policy_b.','Percentage points'),
'lower95_pp':('2.5th percentile of paired exponential household-weight multipliers with policies and fitted scores fixed.','Percentage points'),
'upper95_pp':('97.5th percentile of paired exponential household-weight multipliers with policies and fitted scores fixed.','Percentage points'),
'replications':('Number of paired exponential household-weight replications; intervals are conditional on the selected policies and do not re-estimate their optimum.','Replications'),
'nest':('Consumption nest: upper combines food, energy and other; food, energy and other identify within-nest substitution.','Category'),
'household_group':('Person-weighted benchmark equivalent-income quintile.','Category'),
'elasticity':('Substitution elasticity used in the main nested-CES CGE and household microsimulation.','Dimensionless'),
'source_url':('Public source supporting the calibration reference.','URL'),
'adaptation':('Explicit transfer from the source model to this study; these are calibration assumptions, not Turkish estimates.','Text'),
}

def workbook(name,df):
    assert all(c in META for c in df),set(df)-set(META)
    w=Workbook();s=w.active;s.title='Data';s.append(list(df.columns))
    for row in df.itertuples(index=False,name=None):
        s.append([None if pd.isna(v) else v.item() if isinstance(v,np.generic) else v for v in row])
    d=w.create_sheet('Variable Descriptions');d.append(['Variable','Description','Unit'])
    for c in df:d.append([c,*META[c]])
    p=OUT/'tables/supplementary'/name;w.save(p)
    check=load_workbook(p);assert check.sheetnames==['Data','Variable Descriptions']
    for sheet in check:
        assert sheet.freeze_panes is None and sheet.auto_filter.ref is None
        assert not sheet.merged_cells.ranges and all(not cell.has_style for row in sheet for cell in row)
    return dict(file=p.name,rows=len(df),columns=len(df.columns))

def save(fig,name,supp=False,tr=False):
    folder=(TR/'figures' if tr else OUT/'figures')/('supplementary' if supp else 'main')
    folder.mkdir(parents=True,exist_ok=True)
    if tr:
        fig.canvas.draw()
        for ax in fig.axes:
            for axis in (ax.xaxis,ax.yaxis):
                old=axis.get_major_formatter()
                axis.set_major_formatter(FuncFormatter(lambda x,pos,fmt=old:fmt(x,pos).replace('.',',')))
    for suffix in ['png','svg','pdf','tiff']:
        kw={'pil_kwargs':{'compression':'tiff_lzw'}} if suffix=='tiff' else {}
        fig.savefig(folder/f'{name}.{suffix}',dpi=600,bbox_inches='tight',pad_inches=.04,**kw)
    fd.plt.close(fig)

def framework(tr=False):
    translation=json.loads((TR/'figure_translations.json').read_text(encoding='utf-8')) if tr else {}
    oldsave=fd.save
    def capture(fig,name,caption,supp=False):
        for artist in fig.findobj(Text):
            text=artist.get_text()
            if tr:text=translation.get(text,text)
            if text.startswith(('Household CES elasticity','Household CES esnekliği')):
                text='İç içe CES: gıda · enerji · diğer' if tr else 'Nested CES: food · energy · other'
            if tr and text=='Nested CES: food · energy · other':text='İç içe CES: gıda · enerji · diğer'
            if tr and text=='Own nested-CES cost-of-living':text='Haneye özgü iç içe CES geçinme maliyeti'
            if text.startswith(('Checks and experiments:','Checks ve deneyler:')):
                text=('Politika karşılaştırması → Yararlanıcı bileşimi → PMT kapsamı → Yoksulluk–kapsam sınırı (ε-kısıt)' if tr else
                      'Policy comparison → Recipient composition → PMT coverage → Poverty–coverage frontier (epsilon constraint)')
                artist.set_fontsize(6.1)
            if text.startswith(('Fixed real poverty line','Sabit reel yoksulluk sınırı')):
                text=('Sabit reel yoksulluk sınırı\nFGT0 · FGT1 · FGT2\nTüketim ve\ngelir dağılımı\nGSYH ve yatırım\nScope–yoksulluk\nsınırı' if tr else
                      'Fixed real poverty line\nFGT0 · FGT1 · FGT2\nConsumption and\ninequality\nGDP and investment\nPoverty–coverage\nfrontier')
                artist.set_fontsize(5.9)
            if tr and text.startswith('Household Bütçe Araştırması'):artist.set_fontsize(6.8)
            if tr and text.startswith('CES katma değer'):
                artist.set_linespacing(1.25)
            artist.set_text(text)
        save(fig,'fig1',tr=tr)
    fd.save=capture
    try:fd.fig1_framework()
    finally:fd.save=oldsave

def figures(curve,eps,intervals,robust=None):
    star=curve.loc[curve.fgt1_pct.idxmin()];front=curve[curve.pareto_efficient_evaluated]
    selected=eps[eps.scenario_kind=='retention_target'].sort_values('epsilon_pp')
    for tr in (False,):
        fig,(a,b)=fd.plt.subplots(1,2,figsize=(fd.DOUBLE,3.0),layout='constrained')
        dominated=curve[~curve.pareto_efficient_evaluated]
        a.scatter(dominated.coverage_pct,dominated.fgt1_pct,s=7,color='#C4C7CA',alpha=.8)
        a.plot(front.coverage_pct,front.fgt1_pct,color=fd.COL['pmt'])
        a.scatter([star.coverage_pct],[star.fgt1_pct],marker='*',s=52,color=fd.COL['perfect_targeting'],zorder=5)
        a.annotate(((f'%{star.coverage_pct:.2f}; FGT1 %{star.fgt1_pct:.3f}').replace('.',',') if tr else f'{star.coverage_pct:.2f}%; FGT1 {star.fgt1_pct:.3f}%'),(star.coverage_pct,star.fgt1_pct),xytext=(18,14),textcoords='offset points',fontsize=7,arrowprops={'arrowstyle':'-','color':fd.MUTED,'lw':.6})
        a.set(xlim=(0,102),ylim=(6.60,8.13),xlabel='Nakit alan nüfus (%)' if tr else 'Population receiving cash (%)',ylabel='Yoksulluk açığı, FGT1 (%)' if tr else 'Poverty gap, FGT1 (%)')
        fd.panel(a,'a','Yoksulluk–kapsam sınırı' if tr else 'Poverty–coverage frontier');a.grid(axis='y')
        b.plot(front.fgt1_pct-star.fgt1_pct,front.recipients_million,color=fd.COL['pmt'],lw=1.4)
        absrows=eps[eps.scenario_kind=='absolute_tolerance']
        b.scatter(absrows.epsilon_pp,absrows.recipients_million,s=13,color=fd.COL['pmt'])
        for r in selected.itertuples():
            if r.reduction_retention_requirement_pct==25:continue
            b.scatter(r.epsilon_pp,r.recipients_million,s=27,marker='D',color=fd.COL['universal_cash'],zorder=4)
            pct=int(r.reduction_retention_requirement_pct)
            lab=(f'Kazanımın %{pct}’i\n{r.recipients_million:.2f} milyon'.replace('.',',') if tr else f'{pct}% of gain retained\n{r.recipients_million:.2f} million')
            b.annotate(lab,(r.epsilon_pp,r.recipients_million),xytext=(-27,15),textcoords='offset points',fontsize=6.7)
        b.set(xlim=(0,1.08),ylim=(0,88),xlabel='İzin verilen FGT1 artışı (yüzde puan)' if tr else 'Permitted FGT1 increase (percentage points)',ylabel='Nakit alan persons sayısı (milyon)' if tr else 'Cash recipients (million)')
        fd.panel(b,'b','FGT1 kısıtı altında en geniş kapsam' if tr else 'Maximum coverage under a poverty-gap cap');b.grid(axis='y')
        save(fig,'fig7',tr=tr)
        if tr:continue
        fig,(a,b)=fd.plt.subplots(1,2,figsize=(fd.DOUBLE,3.0),layout='constrained')
        higher=curve[curve.coverage_pct>=star.coverage_pct].sort_values('coverage_pct')
        a.plot(higher.coverage_pct,higher.fgt1_pct-star.fgt1_pct,color=fd.COL['pmt'])
        a.set(xlabel='Nakit alan nüfus (%)' if tr else 'Population receiving cash (%)',ylabel='En düşük FGT1’e göre artış (yp)' if tr else 'Increase from minimum FGT1 (pp)',xlim=(0,101))
        fd.panel(a,'a','Scope genişlemesinin fırsat maliyeti' if tr else 'Opportunity cost of broader coverage');a.grid(axis='y')
        labels=[f'{r.from_coverage_pct:.0f}–{r.to_coverage_pct:.0f}' for r in intervals.itertuples()];labels[0]=f'{intervals.iloc[0].from_coverage_pct:.2f}–{intervals.iloc[0].to_coverage_pct:g}';labels[0]=labels[0].replace('.',',') if tr else labels[0]
        b.barh(np.arange(len(intervals)),intervals.fgt1_pp_per_additional_million,color=fd.COL['pmt'],height=.65)
        b.set(yticks=np.arange(len(labels)),yticklabels=labels,xlabel='Ek 1 milyon persons başına FGT1 artışı (yp)' if tr else 'FGT1 increase per extra million recipients (pp)',ylabel='Scope aralığı (%)' if tr else 'Coverage interval (%)')
        b.invert_yaxis();b.grid(axis='x');b.set_axisbelow(True)
        fd.panel(b,'b','Aralık başına ortalama maliyet' if tr else 'Average incremental cost by interval')
        save(fig,'figS6',supp=True,tr=tr)
    if robust is not None:
        cases=robust.configuration.drop_duplicates().tolist();piv=robust.pivot(index='configuration',columns='choice',values='coverage_pct').reindex(cases)
        lab={'main':'Main','food_only':'Food shock','energy_only':'Energy shock','mild_joint':'Mild joint shock','severe_joint':'Severe joint shock','low_budget':'Budget 0.25% GDP','high_budget':'Budget 1% GDP','low_trade_elasticity':'Trade elasticities ×0.75','high_trade_elasticity':'Trade elasticities ×1.25','low_factor_elasticity':'Factor elasticities ×0.75','high_factor_elasticity':'Factor elasticities ×1.25','zero_admin':'Administration 0%','high_admin':'Administration 10%','energy_support_only':'Energy-support specification','symmetric_world_prices':'Import and export shocks','fixed_negative_investment':'Fixed negative investment','cash_saving_10pct':'Cash saving 10%'}
        for i in range(26,29):lab[f'pmt_partition_202610{i}']=f'PMT partition {i-25}'
        lab.update(consumption_CD='Cobb–Douglas',consumption_flat_CES2='Flat CES: elasticity 2',upper_substitution_050='Upper-nest elasticity 0.50',upper_substitution_085='Upper-nest elasticity 0.85',low_within_substitution='Within-nest elasticities ×0.5',high_within_substitution='Within-nest elasticities ×2')
        fig,ax=fd.plt.subplots(figsize=(fd.DOUBLE,max(4.7,len(cases)*.235)),layout='constrained');y=np.arange(len(cases))
        for col,color,label,marker in [('minimum_gap',fd.COL['perfect_targeting'],'Minimum poverty gap','o'),('retain_75',fd.COL['pmt'],'Retain 75% of maximum reduction','s'),('retain_50',fd.COL['universal_cash'],'Retain 50% of maximum reduction','D')]:
            ax.scatter(piv[col],y,s=18,c=color,label=label,marker=marker)
        ax.set(yticks=y,yticklabels=[lab[c] for c in cases],xlabel='Population receiving cash (%)',xlim=(0,100));ax.invert_yaxis();ax.grid(axis='x');ax.set_axisbelow(True)
        ax.legend(loc='lower right',fontsize=6.5);save(fig,'figS7',supp=True)

def complete_captions(case_count):
    """Keep the caption index consistent with the full set of fourteen figures."""
    path=OUT/'figures/Figure_Captions.md'
    lines=[s for s in path.read_text(encoding='utf-8').splitlines()
           if not any(s.startswith(f'**{key}.**') for key in ['fig7','figS6','figS7'])]
    extra=(' Household demand and household-specific cost-of-living indices use the same '
           'food-energy-other nested-CES structure; consumption-nest elasticities are documented in Table S19.')
    for j,line in enumerate(lines):
        if line.startswith('**fig1.**') and 'consumption-nest elasticities' not in line:
            lines[j]=line+extra
    cap7=('**fig7.** Poverty-coverage frontier and epsilon-constraint choices. '
          '(a) The blue line connects numerically efficient evaluated policies; grey points are dominated among '
          'the evaluated policies, and the star marks the lowest evaluated FGT1. '
          '(b) Maximum cash coverage, expressed as recipient population, under each permitted FGT1 increase '
          'above the numerical minimum. Diamonds retain 75% and 50% of the maximum poverty-gap reduction '
          'relative to the no-policy shock. All points use the same main net fiscal budget, fitted PMT ranking '
          'and equal payment per recipient person; equilibrium prices, incomes and cash envelopes are re-solved.')
    lines.insert(lines.index('## Supplementary figures'),cap7+'\n')
    lines+=['',
        '**figS6.** Opportunity cost of broader PMT coverage. (a) FGT1 increase above its numerical minimum '
        'as coverage expands. (b) The FGT1 increase between each pair of coverage endpoints divided by '
        'the additional recipient population, in percentage points per million additional recipients. '
        'Bars are interval averages. The main net fiscal budget and equal-per-person payment rule are held fixed.',
        '',
        f'**figS7.** Coverage-design robustness across {case_count} economic, preference and PMT-partition cases. '
        'Circles show the numerically refined minimum-gap coverage; squares and diamonds show the largest '
        'coverage retaining 75% and 50% of the maximum poverty-gap reduction. Each row uses its stated '
        'configuration and fiscal envelope, with gains measured relative to that configuration\'s no-policy shock. '
        'The three additional PMT partitions re-estimate out-of-fold scores using alternative fold assignments.',
        '']
    path.write_text('\n'.join(lines),encoding='utf-8')


def main():
    ap=argparse.ArgumentParser();ap.add_argument('--partial',action='store_true');a=ap.parse_args()
    final=json.loads((OUT/'diagnostics/pmt_coverage_optimization_final.json').read_text(encoding='utf-8'))
    eps=pd.read_csv(OUT/'tables/PMT_epsilon_solutions.csv')
    curve=pd.read_csv(OUT/'tables/PMT_pareto_evaluated_policies.csv')
    star=final['optimum'];shock=float(pd.read_csv(OUT/'tables/main/Table1_main_policy_results.csv').set_index('policy').loc['shock','fgt1_pct'])
    curve['fgt1_regret_pp']=curve.fgt1_pct-star['fgt1_pct']
    curve['additional_recipients_million']=curve.recipients_million-star['recipients_million']
    curve['net_cash_billion_tl']=curve.net_cash_billion_tl.fillna(curve.annual_cash_per_recipient_sam_tl*curve.recipients_million/1000)
    curve['maximum_reduction_retained_pct']=100*(shock-curve.fgt1_pct)/(shock-star['fgt1_pct'])
    intervals=pd.read_csv(OUT/'tables/PMT_coverage_tradeoff_intervals.csv')
    summary=None;published=[]
    published.append(workbook('S14_coverage_pareto_frontier.xlsx',curve))
    published.append(workbook('S15_epsilon_constraint_solutions.xlsx',eps))
    if not a.partial:
        audit=json.loads((OUT/'diagnostics/pmt_coverage_robustness_refined.json').read_text(encoding='utf-8'));assert audit['status']=='complete' and audit['completed_case_count']==len(configurations())+3
        summary=pd.read_csv(OUT/'tables/PMT_coverage_robustness_refined_summary.csv')
        rc=pd.read_csv(OUT/'tables/PMT_coverage_robustness_refined_curves.csv');rc['selected_choice']='';rc['selected_retention_pct']=np.nan
        for r in summary.itertuples():
            subset=rc[rc.configuration==r.configuration]
            ix=(subset.coverage_pct-r.coverage_pct).abs().idxmin()
            assert abs(rc.loc[ix,'coverage_pct']-r.coverage_pct)<1e-7
            rc.loc[ix,'selected_choice']=r.choice;rc.loc[ix,'selected_retention_pct']=r.retention_pct
        published.append(workbook('S16_coverage_design_robustness.xlsx',rc))
        published.append(workbook('S17_equilibrium_feedback_decomposition.xlsx',pd.read_csv(OUT/'tables/PMT_coverage_feedback.csv')))
        pairs=pd.DataFrame(final['paired_intervals']).replace({'validated_optimum':'minimum_gap','optimum':'minimum_gap','initial_optimum':'nearby_minimum'})
        published.append(workbook('S18_coverage_conditional_intervals.xlsx',pairs))
        preferences=configurations()['main']['household_nest_sigmas']
        param=[]
        for nest,eta in preferences.items():
            for group in ['Q1','Q2','Q3','Q4','Q5']:
                param.append(dict(nest=nest,household_group=group,elasticity=eta,
                    source_url='https://www3.epa.gov/ttnecas1/models/empax_model_documentation.pdf',
                    adaptation={'upper':'Energy versus other consumption (0.25) adapted to food-energy-other upper nest; also SAGE 1.2.0 Table 9.',
                        'food':'Consumption-goods elasticity (0.5) adapted to A01, A03 and C10-12; not a food-specific estimate.',
                        'energy':'Non-transport energy elasticity (0.4) adapted to C19 and D35, including household transport fuels.',
                        'other':'Consumption-goods elasticity (0.5) adapted to remaining positive-consumption sectors.'}[nest]))
        published.append(workbook('S19_consumption_elasticities.xlsx',pd.DataFrame(param)))
    # Selected article table; the full tolerance menu remains in Table S15.
    idx=(eps.scenario_kind=='retention_target') & eps.reduction_retention_requirement_pct.isin([75,50])
    idx |= (eps.scenario_kind=='absolute_tolerance') & (np.isclose(eps.epsilon_pp,0)|np.isclose(eps.epsilon_pp,.1)|np.isclose(eps.epsilon_pp,.5)|np.isclose(eps.coverage_pct,100))
    short=eps[idx].sort_values('epsilon_pp')
    short[['epsilon_pp','coverage_pct','recipients_million','fgt1_pct','maximum_reduction_retained_pct','target_reached_pct']].to_csv(OUT/'tables/main/table5.csv',index=False,float_format='%.12g')
    figures(curve,eps,intervals,summary);framework(False)
    dump(OUT/'diagnostics/coverage_publication_outputs.json',dict(status='partial' if a.partial else 'complete',tables=published,main_table='table5.csv',main_figure='fig7',supplementary_figures=['figS6','figS7'],metadata=META))
    if not a.partial:
        complete_captions(len(configurations())+3)
        index=OUT/'reports/Supplementary_Tables.md'
        text=index.read_text(encoding='utf-8')
        text='\n'.join(line for line in text.splitlines() if not any(f'[{n}_' in line for n in ['S14','S15','S16','S17','S18','S19']))
        for entry in published:
            topic='Household consumption elasticity calibration' if entry['file'].startswith('S19_') else 'Coverage-design extension'
            text+=f"\n- [{entry['file']}](../tables/supplementary/{entry['file']}): {topic} ({entry['rows']} rows)."
        index.write_text(text.rstrip()+'\n',encoding='utf-8')
    print('Published',len(published),'supplementary tables',flush=True)

if __name__=='__main__':main()
